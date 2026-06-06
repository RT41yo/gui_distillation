#!/usr/bin/env python3
"""Generate deterministic LibreOffice Writer .docx fixtures for synthetic task seeding."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

from ui_explorer.synthetic.paths import repo_root, resolve_repo_path
from ui_explorer.synthetic.writer_fixtures import FIXTURE_FILENAMES, writer_fixtures_dir


def _save(doc: Document, output_dir: Path, filename: str) -> str:
    path = output_dir / filename
    doc.save(path)
    return str(path)


def generate_fixtures(output_dir: Path) -> list[str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    created: list[str] = []

    blank = Document()
    created.append(_save(blank, output_dir, "blank.docx"))

    one_word = Document()
    one_word.add_paragraph("Sample")
    created.append(_save(one_word, output_dir, "one_word.docx"))

    single_sentence = Document()
    single_sentence.add_paragraph("This is a single sentence used for synthetic Writer tasks.")
    created.append(_save(single_sentence, output_dir, "single_sentence.docx"))

    short_paragraph = Document()
    short_paragraph.add_paragraph(
        "This paragraph contains enough text for selection, formatting, and layout tasks."
    )
    created.append(_save(short_paragraph, output_dir, "short_paragraph.docx"))

    two_paragraphs = Document()
    two_paragraphs.add_paragraph("First paragraph for multi-paragraph Writer tasks.")
    two_paragraphs.add_paragraph("Second paragraph with different content.")
    created.append(_save(two_paragraphs, output_dir, "two_paragraphs.docx"))

    bulleted_list = Document()
    bulleted_list.add_paragraph("Alpha item", style="List Bullet")
    bulleted_list.add_paragraph("Beta item", style="List Bullet")
    bulleted_list.add_paragraph("Gamma item", style="List Bullet")
    created.append(_save(bulleted_list, output_dir, "bulleted_list.docx"))

    numbered_list = Document()
    numbered_list.add_paragraph("First numbered item", style="List Number")
    numbered_list.add_paragraph("Second numbered item", style="List Number")
    numbered_list.add_paragraph("Third numbered item", style="List Number")
    created.append(_save(numbered_list, output_dir, "numbered_list.docx"))

    simple_table = Document()
    table = simple_table.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Cell A1"
    table.cell(1, 1).text = "Cell B1"
    created.append(_save(simple_table, output_dir, "simple_table.docx"))

    table_with_values = Document()
    values_table = table_with_values.add_table(rows=3, cols=2)
    values_table.cell(0, 0).text = "Item"
    values_table.cell(0, 1).text = "Qty"
    values_table.cell(1, 0).text = "Apples"
    values_table.cell(1, 1).text = "3"
    values_table.cell(2, 0).text = "Oranges"
    values_table.cell(2, 1).text = "5"
    created.append(_save(table_with_values, output_dir, "table_with_values.docx"))

    styled_heading_and_body = Document()
    heading = styled_heading_and_body.add_paragraph("Project Summary")
    heading.style = "Heading 1"
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    body = styled_heading_and_body.add_paragraph("Body text for style and layout experiments.")
    for run in body.runs:
        run.font.size = Pt(12)
    created.append(_save(styled_heading_and_body, output_dir, "styled_heading_and_body.docx"))

    missing = sorted(set(FIXTURE_FILENAMES) - {Path(path).name for path in created})
    if missing:
        raise RuntimeError(f"fixture generation incomplete, missing: {missing}")
    return created


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=repo_root() / "data/source/fixtures/libreoffice_writer",
        help="Directory where Writer fixture .docx files should be written.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_dir = resolve_repo_path(args.output_dir)
    try:
        created = generate_fixtures(output_dir)
    except ModuleNotFoundError:
        print(json.dumps({
            "ok": False,
            "error": "python-docx is required; install with `pip install python-docx`",
        }, indent=2), file=sys.stderr)
        return 2

    print(json.dumps({
        "ok": True,
        "output_dir": str(output_dir),
        "fixtures": created,
        "count": len(created),
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
